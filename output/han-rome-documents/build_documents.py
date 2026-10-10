from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent

CHINA = {
    'name': '古中国专题_汉代治理与社会生活.docx',
    'title': '汉代中国的治理与社会生活',
    'subtitle': '丝绸之路联系的中国社会背景',
    'intro': '汉代中国参与远距离交换，依靠的是一个具有城市、行政组织、生产活动和边塞设施的社会。理解这些条件，有助于解释为什么向西联系能够持续，也能看见丝绸贸易之外的人与劳动。本文以两汉为中心，通过钱币、简牍和墓葬材料讨论治理与生活，并分析这些材料能怎样补充汉罗联系的研究。',
    'pages': [
        [
            ('一 长安与洛阳的城市背景', [
                '西汉以长安为都，东汉以洛阳为都。大都会艺术博物馆对长安的介绍，将宫殿、居住区域和市场放在同一个城市空间中；对洛阳的介绍，则强调宫殿布局与东汉政治中心的形成。两座都城提供了理解统治、居住和交换的空间背景。[1]',
                '据此可以分析，都城能够汇集官员、使者与物品，但它不能代表整个汉代社会。理解外来珍品在都城出现的意义，还需要追问谁能够接触这些物品，以及城市之外的生产与运输如何支持这种集中。'
            ]),
            ('二 经典教育与行政秩序', [
                '汉武帝时期，儒家经典在官方行为规范和教育中取得重要地位。《史记》也为后来的历史书写提供了重要范式。经典与史书共同影响了官员理解政治秩序及周边人群的方式。[1]',
                '思想规范与实际行政应分开考察。官方倡导某种道德，并不能说明每位官员都按同一种方式行事；一部史书用汉朝立场描述远方，也不能说明被描述的社会接受相同的分类。研究汉代对西域的认识，需要同时注意观察内容与叙述立场。'
            ]),
            ('三 居延汉简中的边塞日常', [
                '居延位于汉代西北边境。中央研究院历史语言研究所的藏品介绍指出，汉武帝以来，边地的移民、屯田、烽燧、邮驿及道路共同构成防卫体系。边塞人口包括来自内郡的戍卒、田卒等不同群体。[2]',
                '相关简牍保存了谷物出入、人员登记和邮件传递等记录。这些材料使边塞历史落实到供粮、值守与文书事务。由此可以提出一个解释：远方联系能否维持，也取决于沿途人员能否持续获得补给、完成组织工作。这一解释不意味着国家邮驿对所有商人自由开放。[2]'
            ])
        ],
        [
            ('四 五铢钱与铸币管理', [
                '中国国家博物馆介绍的上林三官五铢，始铸于汉武帝元鼎四年，即公元前113年。上林三官分担铸造、铜料整理和刻范等职责，币制改革把铸币权与铸造活动集中到中央。这是理解汉代国家与货币关系的一个具体案例。[3]',
                '从经济组织角度看，货币制度使研究者能够追问标准、发行与交易的关系。但同类钱币在不同地点出现，不足以证明所有地区的市场都一样，也不足以确定某条跨国商路。钱币需要与出土年代、同出物及当地使用环境一起分析。'
            ]),
            ('五 盐场画像砖中的生产环节', [
                '国博的东汉盐场画像砖呈现蜀地井盐生产：井架上的劳动者提取卤水，枧筒把卤水送往煮盐处，灶前有人烧火，另有人背负盐包运输。图像把取卤、煮盐与搬运安排在同一场景中，体现了生产过程的衔接。[4]',
                '这个案例提醒我们，汉代经济还需要从区域资源和劳动分工来理解。盐场材料可以帮助解释生产组织，却不能自动证明画中盐品进入丝绸之路，更不能把它改写为汉朝对罗马的出口记录。生产能力与远距离销售是两个需要分别取证的问题。'
            ]),
            ('六 明器所呈现的生活与愿望', [
                '汉代墓葬中的明器包括房屋、粮仓、井、灶、牲畜栏以及人物模型。密歇根大学艺术博物馆介绍的多层建筑模型，属于为墓葬制作的生活环境缩影。这些物品保存了关于居所、储藏与服务活动的线索。[5]',
                '墓葬材料经过选择，往往也表达家属对死者生活的安排与期待。因此，一座模型可以提供建筑与社会观念的证据，却不是全国住房的统计样本。将明器与边塞简牍对读，能够比较被纪念的生活和实际记账的生活，同时保留地域与人群差异。'
            ])
        ],
        [
            ('七 汉代社会如何认识远方', [
                '汉文资料中的大秦通常对应罗马帝国或其控制地区，但具体范围应随语境判断。关于远方的知识，可能经过商人、使者及后来的编纂者传递。地名解释与文明描述不能一律当作亲历观察。[6]',
                '前述社会材料提供了进一步的分析角度。负责边塞文书的人、生产盐的劳动者和使用墓葬明器的家庭，各自接触远方消息的条件可能不同。一个帝国的官方记录，不等于其所有居民拥有同一幅世界地图。这是根据材料性质提出的解释，不能据此编造具体人物的知识与经历。'
            ]),
            ('八 将社会背景用于汉罗联系研究', [
                '分析一件外来物品时，可以依次核对其生产条件、流通环境与使用场景。例如，货币管理解释的是交易制度，简牍揭示的是组织事务，明器呈现的是被选择的生活形象。这些证据互相补充，但不能代替有关进口物产或使节来访的专门材料。',
                '值得继续讨论的问题是：边塞补给怎样影响长途联系？统一铸币如何与地域市场共存？墓葬珍品怎样表达家庭地位？官方文献和日常文书对远方的关注为何可能不同？这些问题将汉罗联系放回汉代社会之中，而不是把整个社会缩减为丝绸生产者。'
            ])
        ]
    ],
    'refs': [
        ('大都会艺术博物馆', 'Han Dynasty 206 BC to 220 AD', 'https://www.metmuseum.org/essays/han-dynasty-206-b-c-220-a-d'),
        ('中央研究院历史语言研究所历史文物陈列馆', '居延汉简', 'https://museum.sinica.edu.tw/collection/20/'),
        ('中国国家博物馆', '五铢铜钱', 'https://www.chnmuseum.cn/zp/zpml/hb/201812/t20181218_26289.shtml'),
        ('中国国家博物馆', '盐场画像砖', 'https://www.chnmuseum.cn/zp/zpml/kgfjp/202110/t20211028_251933.shtml'),
        ('密歇根大学艺术博物馆', 'Model of a four storied pavilion', 'https://umma.umich.edu/objects/model-of-a-four-storied-pavilion-1993-1-71-3/'),
        ('John E Hill 华盛顿大学 Silk Road Seattle', '后汉书大秦条地名注释', 'https://depts.washington.edu/silkroad/texts/hhshu/notes11.html'),
    ]
}

ROME = {
    'name': '古罗马专题_城市消费与东方贸易.docx',
    'title': '古罗马的城市消费与东方贸易',
    'subtitle': '丝绸之路联系的罗马社会背景',
    'intro': '罗马世界对东方商品的需求，需要放在城市生活、工艺生产和社会观念中理解。进口商品既涉及远方交换，也进入本地的使用与评价。本文以公元前1世纪至公元2世纪为中心，从港口记述、玻璃器、军人书信和哲学文本展开，说明罗马社会为什么需要这些商品，以及不同材料怎样呈现彼此不完全相同的生活。',
    'pages': [
        [
            ('一 皇帝权力与帝国中的不同社会', [
                '公元前27年，屋大维获得奥古斯都称号。大都会艺术博物馆的藏品说明指出，他的统治将共和国传统与实际的君主权力结合起来，权力通过职务、特权及军队控制来维持。因此，早期帝国的制度不能只凭“皇帝”一词来理解。[1]',
                '帝国内部也存在地方差异。大英博物馆以罗马不列颠为例，说明罗马的行政、货币与建筑等因素和当地社会相遇，形成具有地域特点的生活方式。研究消费时，应区分罗马城、东方港口和不同边疆，避免把首都经验视为所有居民的经验。[2]'
            ]),
            ('二 印度港口中的经营知识', [
                '约公元1世纪的《厄立特里亚海航行记》是一部有关红海和印度洋贸易的古代记述。其第43至46节讨论Barygaza港口附近的浅滩、潮汐和当地引航。第49节列出输入的葡萄酒、金属等商品，以及输出的棉织品、丝织品等货物。[3]',
                '这些记述显示，跨海经营需要当地知识与服务。船只到达一个海湾，并不等于能安全进港；货物列在港口清单上，也不等于它在该港生产。这些细节可用于解释罗马商人如何接触东方商品，而不必预设所有商人都亲自到达商品的最初产地。'
            ]),
            ('三 玻璃在日常生活中的用途', [
                '罗马玻璃可用于饮食器皿、香油及药物容器，也用于建筑和装饰。大都会艺术博物馆介绍了玻璃吹制带来的器形与生产变化；吹制发展之后，铸制工艺仍然继续存在。因此，“罗马玻璃”包含不同工艺和用途。[4]',
                '从使用场景来看，玻璃并不只有珍贵礼品这一种身份。同一材料可以成为日用品，也可以经过复杂加工成为昂贵器物。判断一件玻璃器的社会意义，应考察具体器形、制作方式和使用环境，而不是只依据材料名称。'
            ])
        ],
        [
            ('四 套色浮雕玻璃与高价器物', [
                '罗马套色浮雕玻璃需要制作多层玻璃，再对外层进行精细雕刻。大都会艺术博物馆指出，这类器物生产复杂、耗时而昂贵，早期帝国的皇室和元老精英是重要的消费群体。多数相关遗存属于奥古斯都及朱里亚克劳狄时期。[5]',
                '这个案例有助于分析，商品价值可以来自加工难度与展示用途，而不只是遥远产地。进一步讨论东方丝绸时，也应询问织物怎样被使用、改制或赋予地位意义。不过，套色玻璃与丝绸的消费机制只能作为比较，不能把一个案例的结论直接套用于另一种商品。'
            ]),
            ('五 军人书信所呈现的生活', [
                '大英博物馆的罗马军队展览以克劳狄乌斯特伦提亚努斯的经历，以及文德兰达书写板等材料，呈现军旅生活。相关资料还涉及与军队共同生活的女性、儿童和被奴役者。军队因此也是研究人员移动与生活关系的入口。[6]',
                '把这些材料与精美器物对读，可以避免只用皇帝或富裕家庭的消费来定义罗马生活。书信与书写板显示具体交往，却同样具有保存和识字人群的偏向。它们不能自动代表全部士兵，更不能在缺少内容支持时被改写成军队向中国运送商品的记录。'
            ]),
            ('六 消费批评与斯多葛哲学', [
                '老普林尼在《自然史》12.29至30附近讨论胡椒与远方商品时，批评人们为口腹和香气远行求取物产。这种文字既保存商品信息，也表达作者对欲望的评价。读者需要区分“有人消费某物”和“作者如何评价消费”。[7]',
                '马可奥勒留的《沉思录》则是理解斯多葛实践的重要文本。斯坦福哲学百科指出，他把德性与外在财富等条件区分开来。这种思想材料不能代表所有罗马消费者，也没有因此证明其与汉代思想建立直接传承。它的价值在于提供另一种评价生活与欲望的语言。[8]'
            ])
        ],
        [
            ('七 认识商品与认识产地之间的距离', [
                '港口经营者关注航行条件与可交易货物，工匠关注材料与制作，哲学作者关注行为与价值。不同活动产生不同的知识。一个人熟悉丝织品的触感与价格，并不因此知道遥远产地的行政制度；一个文本描述外来物产，也未必提供了可靠的整段运输路线。',
                '因此，研究罗马与汉代中国的联系，可先问某条证据属于哪一种活动。贸易清单支持有关港口交换的判断，工艺遗存支持制作与使用研究，哲学文本支持观念分析。只有将这些证据分别定位，才有可能重建它们之间的联系，而不是把所有信息合成一幅没有差异的帝国图景。'
            ]),
            ('八 将罗马社会材料用于跨区域比较', [
                '罗马玻璃与汉代明器，可以比较器物怎样参与生活与地位表达；罗马军人书写材料与居延汉简，可以比较边疆中的人员、供给和文书。这样的比较需要尊重材料类型：随葬模型、经营记述和日常书信各自回答不同的问题。',
                '进一步的研究问题包括：港口引航怎样影响贸易组织？同种材料为何同时进入日常与精细消费？作者对奢侈的批评与实际需求为何可能并存？边疆材料怎样补充首都叙事？比较能提出解释，但相似之处本身不证明两地制度或思想之间存在直接借用。'
            ])
        ]
    ],
    'refs': [
        ('大都会艺术博物馆', 'Colossal marble head of the emperor Augustus', 'https://www.metmuseum.org/art/collection/search/251114'),
        ('大英博物馆', 'Roman Britain', 'https://www.britishmuseum.org/collection/galleries/roman-britain'),
        ('匿名古代作者 华盛顿大学 Silk Road Seattle 英译', '厄立特里亚海航行记 第43至49节', 'https://depts.washington.edu/silkroad/texts/periplus/periplus.html'),
        ('Rosemarie Trentinella 大都会艺术博物馆', 'Roman Glass', 'https://www.metmuseum.org/essays/roman-glass'),
        ('Rosemarie Trentinella 大都会艺术博物馆', 'Roman Cameo Glass', 'https://www.metmuseum.org/fr/essays/roman-cameo-glass'),
        ('大英博物馆', 'Legion Life in the Roman army', 'https://www.britishmuseum.org/exhibitions/legion-life-roman-army'),
        ('老普林尼 Attalus 英译', '自然史 第12卷 第29至30节', 'https://www.attalus.org/pliny/hn12a.html'),
        ('斯坦福哲学百科', 'Marcus Aurelius', 'https://plato.stanford.edu/entries/marcus-aurelius/'),
    ]
}


def set_font(style, size, east='Songti SC', bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for key, value in [('ascii', 'Times New Roman'), ('hAnsi', 'Times New Roman'), ('eastAsia', east)]:
        fonts.set(qn('w:' + key), value)
    for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        fonts.attrib.pop(qn('w:' + key), None)


def link(paragraph, url):
    rel = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    elem = OxmlElement('w:hyperlink')
    elem.set(qn('r:id'), rel)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '333333'); props.append(color)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '18'); props.append(size)
    run.append(props)
    text = OxmlElement('w:t'); text.text = url; run.append(text)
    elem.append(run); paragraph._p.append(elem)


def build(data):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(.78); sec.bottom_margin = Inches(.75)
    sec.left_margin = Inches(.82); sec.right_margin = Inches(.82)
    sec.footer_distance = Inches(.34)
    normal = doc.styles['Normal']
    set_font(normal, 11)
    normal.paragraph_format.line_spacing = Pt(17)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    for name, size in [('Title', 23), ('Subtitle', 11), ('Heading 1', 14), ('Heading 2', 12)]:
        st = doc.styles[name]; set_font(st, size, bold=(name != 'Subtitle'))
        st.font.italic = False
        st.paragraph_format.line_spacing = Pt(30 if name == 'Title' else 22 if name == 'Heading 1' else 17)
        st.paragraph_format.space_before = Pt(12 if name == 'Heading 1' else 0)
        st.paragraph_format.space_after = Pt(7 if name != 'Title' else 9)
        st.paragraph_format.keep_with_next = True
        borders = st.element.get_or_add_pPr().find(qn('w:pBdr'))
        if borders is not None: st.element.get_or_add_pPr().remove(borders)
    for name in ['Header', 'Footer']:
        set_font(doc.styles[name], 9)
    ref_style = doc.styles.add_style('Source entry', 1)
    ref_style.base_style = normal
    set_font(ref_style, 9)
    ref_style.paragraph_format.line_spacing = Pt(12)
    ref_style.paragraph_format.space_after = Pt(5)
    ref_style.paragraph_format.keep_together = True
    ref_style.paragraph_format.keep_with_next = False
    doc.add_paragraph(data['title'], 'Title')
    doc.add_paragraph(data['subtitle'], 'Subtitle')
    meta = doc.add_paragraph('历史专题阅读资料  ·  2026年10月8日')
    meta.paragraph_format.space_after = Pt(12)
    meta.runs[0].font.size = Pt(9)
    doc.add_paragraph(data['intro'])
    for i, blocks in enumerate(data['pages']):
        for title, paragraphs in blocks:
            doc.add_paragraph(title, 'Heading 1')
            for text in paragraphs:
                p = doc.add_paragraph(text)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph('参考资料', 'Heading 1')
    p = doc.add_paragraph('正文方括号对应以下来源。史实陈述与基于材料提出的解释分别表述。资料检索日期为2026年10月8日。', 'Source entry')
    for i, (author, title, url) in enumerate(data['refs'], 1):
        p = doc.add_paragraph(f'[{i}] {author}  {title}', 'Source entry')
        p.add_run('\n')
        link(p, url)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run()
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    fr = OxmlElement('w:r'); ft = OxmlElement('w:t'); ft.text = '1'
    fr.append(ft); field.append(fr); r._r.addnext(field)
    doc.core_properties.title = data['title']
    doc.core_properties.subject = data['subtitle']
    doc.core_properties.author = 'Codex'
    doc.core_properties.keywords = '古代历史 汉代中国 罗马帝国 丝绸之路 社会生活'
    doc.core_properties.comments = ''
    path = ROOT / data['name']
    doc.save(path)
    print(path)


if __name__ == '__main__':
    for item in [CHINA, ROME]:
        build(item)
